import os
import glob
import csv
from xlsxwriter.workbook import Workbook
from datetime import datetime

import openpyxl



def report_gen(template_name,sheet_name,chart_cell_name):
        current_time = datetime.now().strftime("%Y%m%d%H%M%S")
        csvfilename = template_name+".csv"
        templatename = template_name+".xlsx"
        outputfilename = template_name+"_"+current_time+".xlsx"

        #open required worksheet tab in template excel and get coordinates of "Pie" string from first header row
        template_sheet = openpyxl.load_workbook(templatename)[sheet_name]
        i = 0
        for cell in template_sheet[1]:
            i = i+1
            if(cell.value == chart_cell_name):
                pie_chart_coordinate = i

        print(pie_chart_coordinate)

        #copy csv file data from csv to target excel sheet
        workbook = Workbook(outputfilename)
        item_qty_fmt = workbook.add_format({'num_format': '#,##0'})

        worksheet = workbook.add_worksheet(sheet_name)


        with open(csvfilename, 'rt', encoding='utf8') as f:
            reader = csv.reader(f)
            #headers = next(reader, None)
            for r, row in enumerate(reader):
                for c, col in enumerate(row):
                    #skip first line formatting cause its header
                    #format 2nd column to numeric format
                    if (r != 0) and (c == 1):
                        #print(r, c, "format loop", col)
                        worksheet.write_number(r, c, int(col), item_qty_fmt)
                    else:
                        #print(r, c)
                        worksheet.write(r,c,col)
        #print("Total rows are {rows} and total columns are {cols}".format(rows=r, cols=c))


        chart_cordinates_a = sheet_name+"!$A$2:$A$"+str(r+1)
        chart_cordinates_b = sheet_name+"!$B$2:$B$"+str(r+1)

        #print(chart_cordinates_a, chart_cordinates_b)
        #now create pie chart and place it in the same worksheet, at pie chart coordinate
        pie_chart = workbook.add_chart({'type': 'pie'})
        pie_chart.add_series({
            'name': 'Pie data',
            'categories': chart_cordinates_a,
            'values': chart_cordinates_b
        })
        worksheet.insert_chart(row=1, col=pie_chart_coordinate, chart=pie_chart)
        workbook.close()



if __name__ == "__main__":
    #write code to accept template name as argument
    #or path to the template file, and template name
    report_gen("data","SheetName","Pie")
